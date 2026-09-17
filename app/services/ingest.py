"""Turn a URL or an uploaded file into plain text plus provenance.

The extraction interface is final; the extractors themselves are deliberately
simple. Day 1's pipeline reuses `extract_html` for the Jenosize corpus scrape,
so improvements there benefit both paths.
"""

from __future__ import annotations

import hashlib
import io
from typing import Final
from uuid import UUID, uuid4

import anyio
import httpx

from app.core.errors import UpstreamError, ValidationError
from app.core.logging import get_logger
from app.schemas.sources import SourceDocument
from app.storage import Storage
from app.storage.keys import raw_scrape_key, upload_key

logger = get_logger(__name__)

# A browser-ish UA: several publisher CDNs return 403 to the default httpx UA.
_UA: Final = "Mozilla/5.0 (compatible; JenosizeTrendWriter/0.1; +https://trend-writer.workser.app)"
_FETCH_TIMEOUT_S: Final = 20.0
_MAX_FETCH_BYTES: Final = 5 * 1024 * 1024

SUPPORTED_UPLOAD_SUFFIXES: Final = (".pdf", ".docx", ".txt", ".md")


def content_hash(data: bytes) -> str:
    """Stable id for deduplicating scrapes and uploads (see the UNIQUE column)."""
    return hashlib.sha256(data).hexdigest()


def extract_html(html: str, *, url: str | None = None) -> tuple[str | None, str]:
    """HTML -> (title, main text). Boilerplate, nav and comments are dropped."""
    import trafilatura
    from trafilatura.metadata import extract_metadata

    text = trafilatura.extract(html, url=url, include_comments=False, include_tables=True) or ""
    title = None
    try:
        if meta := extract_metadata(html):
            title = meta.title
    except Exception:
        logger.warning("metadata_extract_failed", extra={"url": url})
    return title, text.strip()


def extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()


def extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    return "\n\n".join(p.text.strip() for p in document.paragraphs if p.text.strip())


def extract_plaintext(data: bytes) -> str:
    return data.decode("utf-8", errors="replace").strip()


async def ingest_url(url: str, storage: Storage | None = None) -> SourceDocument:
    """Fetch a page, archive the raw HTML, and return its extracted text."""
    try:
        async with httpx.AsyncClient(
            timeout=_FETCH_TIMEOUT_S,
            follow_redirects=True,
            headers={"User-Agent": _UA},
        ) as client:
            response = await client.get(url)
            response.raise_for_status()
    except httpx.HTTPError as exc:
        raise UpstreamError(f"Could not fetch {url}: {exc}") from exc

    raw = response.content[:_MAX_FETCH_BYTES]
    digest = content_hash(raw)

    # Extraction is CPU-bound lxml work; keep it off the event loop.
    title, text = await anyio.to_thread.run_sync(
        lambda: extract_html(raw.decode(response.encoding or "utf-8", errors="replace"), url=url)
    )

    r2_key = None
    if storage is not None:
        # Archive the raw HTML: re-extraction with a better parser later should
        # not require re-crawling (and the publisher may have changed the page).
        r2_key = await storage.put_bytes(raw_scrape_key(digest), raw, content_type="text/html")

    return SourceDocument(
        id=uuid4(),
        kind="url",
        title=title or url,
        url=url,
        r2_key=r2_key,
        text=text,
        text_chars=len(text),
    )


async def ingest_upload(
    filename: str,
    data: bytes,
    storage: Storage | None = None,
    *,
    source_id: UUID | None = None,
) -> SourceDocument:
    """Extract text from an uploaded pdf/docx/txt/md and archive the original."""
    lower = filename.lower()
    if not lower.endswith(SUPPORTED_UPLOAD_SUFFIXES):
        raise ValidationError(
            f"Unsupported file type. Supported: {', '.join(SUPPORTED_UPLOAD_SUFFIXES)}"
        )

    sid = source_id or uuid4()

    def _extract() -> str:
        if lower.endswith(".pdf"):
            return extract_pdf(data)
        if lower.endswith(".docx"):
            return extract_docx(data)
        return extract_plaintext(data)

    try:
        text = await anyio.to_thread.run_sync(_extract)
    except Exception as exc:
        raise ValidationError(f"Could not read {filename}: {exc}") from exc

    r2_key = None
    if storage is not None:
        r2_key = await storage.put_bytes(upload_key(sid, filename), data)

    return SourceDocument(
        id=sid,
        kind="upload",
        title=filename,
        r2_key=r2_key,
        text=text,
        text_chars=len(text),
    )
